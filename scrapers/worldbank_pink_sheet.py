"""
World Bank Pink Sheet Precious Metal Prices
===========================================

Reads monthly gold, platinum and silver prices from the World Bank's
Commodity Price Data ("Pink Sheet") workbook. The World Bank republishes it
early each month with the previous month's averages, back to 1960:

    https://www.worldbank.org/en/research/commodity-markets
        -> "Monthly prices" -> CMO-Historical-Data-Monthly.xlsx

The workbook's URL holds a hash that changes with each release, so it is
read from the landing page every run.

  1. GET the landing page and take the one link to
     ``CMO-Historical-Data-Monthly.xlsx``.
  2. GET the workbook. On its "Monthly Prices" sheet, find the header row
     (the one with "Gold" in it) and check that the row under it gives every
     wanted metal's unit as ``($/troy oz)``. Each later row is one month,
     ``1960M01`` in the first column.

Output is tidy / long, one row per metal and month:

    metal      — ``gold`` | ``platinum`` | ``silver``
    commodity  — ``Gold`` | ``Platinum`` | ``Silver``
    market     — ``Monthly average``
    unit       — ``USD/oz`` (troy ounce)
    date       — ISO date of the month's 1st (YYYY-MM-01)
    price      — float, the month's average in nominal US dollars

Suggested job key: ``[commodity, market, date]``. The commodity and market
labels are fixed here rather than read from the workbook, because the server
builds series codes from them: a renamed column or unit must fail the run,
not start new series. Months with a blank cell (or ``…``) have no row.

Parameters (job YAML ``parameters``):
    metals         only these metals, by slug     (default: all three)
    delay_min/max  seconds between requests when the job has no ``rate_limit``

The workbook is the whole history in one file, so any failure fails the
run; there are no partial marker rows.
"""

import io
import logging
import re
from typing import Any, Dict, List, Optional

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from scrapers.scrap_monster_prices import _Pacer, _USER_AGENT
from utils.http import get_with_retries

logger = logging.getLogger(__name__)

LANDING_URL = "https://www.worldbank.org/en/research/commodity-markets"
WORKBOOK_NAME = "CMO-Historical-Data-Monthly.xlsx"
SHEET = "Monthly Prices"

# metal slug -> column header in the workbook, which is also the commodity label
METALS = {
    "gold": "Gold",
    "platinum": "Platinum",
    "silver": "Silver",
}
MARKET = "Monthly average"
SHEET_UNIT = "($/troy oz)"
UNIT = "USD/oz"

# Cells the World Bank leaves for a month without a price.
_BLANKS = ("", "…", "...")
_PERIOD_RE = re.compile(r"^(\d{4})M(\d{2})$")


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def find_workbook_url(html: str) -> str:
    """The landing page's link to the monthly workbook. Raises ValueError
    unless there is exactly one."""
    soup = BeautifulSoup(html, "html.parser")
    hrefs = {a["href"].strip() for a in soup.select("a[href]")
             if a["href"].strip().split("?")[0].endswith("/" + WORKBOOK_NAME)}
    if len(hrefs) != 1:
        raise ValueError(f"expected one link to {WORKBOOK_NAME} on {LANDING_URL}, found {sorted(hrefs)}")
    return hrefs.pop()


def parse_period(text: Any) -> str:
    """``"1960M01"`` -> ``"1960-01-01"``. Raises ValueError otherwise."""
    match = _PERIOD_RE.match(str(text).strip())
    if not match or not 1 <= int(match.group(2)) <= 12:
        raise ValueError(f"unexpected period {text!r}")
    return f"{match.group(1)}-{match.group(2)}-01"


def parse_price(value: Any) -> Optional[float]:
    """A number -> float; a blank cell or ``…`` -> None. Raises ValueError
    for any other text, so a stray note never becomes a price."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if str(value).strip() in _BLANKS:
        return None
    raise ValueError(f"unexpected price cell {value!r}")


def parse_workbook(content: bytes, metals: List[str]) -> List[Dict[str, Any]]:
    """Rows for the wanted metals from the workbook's bytes.

    Raises ValueError if the sheet, a metal's column or its ``($/troy oz)``
    unit is missing, or a month's period or price does not parse.
    """
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise ValueError(f"no {SHEET!r} sheet in the workbook (sheets: {book.sheetnames})")
        table = list(book[SHEET].iter_rows(values_only=True))
    finally:
        book.close()

    header_at = next((i for i, row in enumerate(table)
                      if any(str(c).strip() == "Gold" for c in row if c is not None)), None)
    if header_at is None or header_at + 1 >= len(table):
        raise ValueError(f"no header row with 'Gold' on the {SHEET!r} sheet")
    headers = [str(c).strip() if c is not None else "" for c in table[header_at]]
    units = [str(c).strip() if c is not None else "" for c in table[header_at + 1]]

    columns = {}
    for metal in metals:
        commodity = METALS[metal]
        if headers.count(commodity) != 1:
            raise ValueError(f"expected one {commodity!r} column on the {SHEET!r} sheet")
        col = headers.index(commodity)
        if units[col] != SHEET_UNIT:
            raise ValueError(f"{commodity} is quoted in {units[col]!r}, not {SHEET_UNIT!r}")
        columns[metal] = col

    rows: List[Dict[str, Any]] = []
    for row in table[header_at + 2:]:
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        date = parse_period(row[0])
        for metal, col in columns.items():
            price = parse_price(row[col] if col < len(row) else None)
            if price is not None:
                rows.append({"metal": metal, "commodity": METALS[metal], "market": MARKET,
                             "unit": UNIT, "date": date, "price": price})
    return rows


# ---------------------------------------------------------------------------
# Fetching
# ---------------------------------------------------------------------------


class _Fetcher:
    def __init__(self, pacer: _Pacer):
        self.pacer = pacer

    def get(self, url: str) -> bytes:
        self.pacer.wait()
        logger.info("GET %s", url)
        response = get_with_retries(url, headers={"User-Agent": _USER_AGENT}, timeout=120)
        response.raise_for_status()
        return response.content


def run(params: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """Returns tidy rows ``{metal, commodity, market, unit, date, price}``."""
    params = params or {}
    metals = [str(m).lower() for m in (params.get("metals") or METALS)]
    unknown = [m for m in metals if m not in METALS]
    if unknown:
        raise ValueError(f"unknown metals {unknown}; the Pink Sheet has {list(METALS)}")

    fetcher = _Fetcher(_Pacer(float(params.get("delay_min", 1.0)),
                              float(params.get("delay_max", 2.0))))
    landing = fetcher.get(LANDING_URL).decode("utf-8", errors="replace")
    rows = parse_workbook(fetcher.get(find_workbook_url(landing)), metals)
    for metal in metals:
        dates = [r["date"] for r in rows if r["metal"] == metal]
        if not dates:
            raise ValueError(f"no {metal} prices in the workbook")
        logger.info("%s: %d months, %s to %s", METALS[metal], len(dates), min(dates), max(dates))
    return rows
